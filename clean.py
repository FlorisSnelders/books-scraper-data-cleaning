"""
Clean the raw scrape and deliver it in client-ready formats.

Reads data/raw_books.csv (written by scrape.py) and produces:
  - data/clean_books.csv    clean data, one row per book
  - data/clean_books.xlsx   the same data, formatted for Excel
  - cleaning_report.md      what was wrong, what was fixed, before/after examples

Run:  python clean.py
"""

import logging
import re
from collections.abc import Callable
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
RAW_CSV = ROOT / "data" / "raw_books.csv"
CLEAN_CSV = ROOT / "data" / "clean_books.csv"
CLEAN_XLSX = ROOT / "data" / "clean_books.xlsx"
REPORT_MD = ROOT / "cleaning_report.md"

CURRENCY_CODES = {"£": "GBP", "$": "USD", "€": "EUR"}
RATING_WORDS = {"one": 1, "two": 2, "three": 3, "four": 4, "five": 5}
COLUMN_ORDER = ["title", "category", "price", "currency", "star_rating", "availability", "upc", "product_page_url"]

log = logging.getLogger("clean")


# --- Cleaning functions: one job each, covered by tests/test_clean.py ---

def fix_encoding(text: str) -> str:
    """Repair UTF-8 text that was decoded as Latin-1, e.g. "Â£" -> "£".

    Text that doesn't have this problem is returned unchanged.
    """
    try:
        return text.encode("latin-1").decode("utf-8")
    except UnicodeError:
        return text


def normalize_whitespace(text: str) -> str:
    """Trim both ends and collapse runs of spaces, tabs and newlines into one space."""
    return re.sub(r"\s+", " ", text).strip()


def to_snake_case(name: str) -> str:
    """Column name to snake_case: "Product Page URL" -> "product_page_url"."""
    return re.sub(r"[^0-9a-zA-Z]+", "_", name).strip("_").lower()


def parse_price(text: str) -> float | None:
    """Number from a price string: "£1,299.00" -> 1299.0 (None if there is no number)."""
    match = re.search(r"\d[\d,]*(\.\d+)?", text)
    return float(match.group().replace(",", "")) if match else None


def parse_currency(text: str) -> str | None:
    """Currency code from a price string: "£51.77" -> "GBP" (None if unknown)."""
    for symbol, code in CURRENCY_CODES.items():
        if symbol in text:
            return code
    return None


def parse_rating(text: str) -> int | None:
    """Star rating word to integer: "Three" -> 3 (None if it isn't One to Five)."""
    return RATING_WORDS.get(text.strip().lower())


def parse_availability(text: str) -> int | None:
    """Copies in stock: "In stock (22 available)" -> 22, "Out of stock" -> 0.

    Returns None if the text doesn't say how many are available.
    """
    match = re.search(r"\((\d+) available\)", text)
    if match:
        return int(match.group(1))
    if "out of stock" in text.lower():
        return 0
    return None


# --- The cleaning pipeline ---

def as_code(value, max_length: int = 60) -> str:
    """Format a value as Markdown code. Strings use repr() so hidden
    whitespace (\\n) and broken characters (\\x80) stay visible."""
    text = repr(value) if isinstance(value, str) else str(value)
    if len(text) > max_length:
        text = text[:max_length - 3] + "..."
    return "`" + text.replace("|", "\\|") + "`"


def apply_fix(df: pd.DataFrame, issues: list[dict], column: str, func: Callable,
              description: str, target: str | None = None) -> None:
    """Apply `func` to every value in `column` (writing to `target` if given)
    and record how many rows it actually changed, with one real example."""
    before = df[column]
    after = before.map(func)
    changed = after.notna() & (before != after)
    if changed.any():
        first = changed.idxmax()
        issues.append({
            "issue": description,
            "rows": int(changed.sum()),
            "example": f"{as_code(before[first])} → {as_code(after[first])}",
        })
        log.info("%-50s %4d rows", description, changed.sum())
    df[target or column] = after


def drop_rows(df: pd.DataFrame, issues: list[dict], mask: pd.Series, description: str) -> pd.DataFrame:
    """Remove the rows where `mask` is True and record how many were removed."""
    log.info("%-50s %4d rows", description, mask.sum())
    if mask.any():
        first = mask.idxmax()
        issues.append({
            "issue": description,
            "rows": int(mask.sum()),
            "example": f"UPC {as_code(df.at[first, 'upc'])}",
        })
    return df[~mask]


def validate(df: pd.DataFrame) -> list[tuple[str, bool]]:
    """Sanity checks the clean data must pass before it is delivered."""
    return [
        ("Data is not empty", len(df) > 0),
        ("No missing values", bool(df.notna().all().all())),
        ("No duplicate rows", not df.duplicated().any()),
        ("Every UPC is unique", df["upc"].is_unique),
        ("Every price is above 0", bool((df["price"] > 0).all())),
        ("Every star_rating is between 1 and 5", bool(df["star_rating"].between(1, 5).all())),
        ("Every availability is 0 or more", bool((df["availability"] >= 0).all())),
    ]


def clean_books(raw: pd.DataFrame) -> tuple[pd.DataFrame, list[dict]]:
    """Run every cleaning step on the raw data.

    Returns the clean data and the list of issues that were found and fixed.
    Raises ValueError if the result fails validation, so bad data is never
    delivered silently.
    """
    df = raw.copy()
    issues: list[dict] = []

    # 1. Consistent column names: "Star Rating" -> "star_rating"
    df.columns = [to_snake_case(c) for c in df.columns]

    # 2. Repair broken characters, then stray whitespace, in every column
    for column in df.columns:
        apply_fix(df, issues, column, fix_encoding, f"Encoding artifacts in `{column}`")
    for column in df.columns:
        apply_fix(df, issues, column, normalize_whitespace, f"Stray whitespace/newlines in `{column}`")

    # 3. Convert text to proper types (currency first, before price becomes a number)
    apply_fix(df, issues, "price", parse_currency, "Currency symbol split into `currency` column", target="currency")
    apply_fix(df, issues, "price", parse_price, "`price` converted from text to number")
    apply_fix(df, issues, "star_rating", parse_rating, "`star_rating` converted from word to 1-5")
    apply_fix(df, issues, "availability", parse_availability, "`availability` converted to number in stock")

    # 4. Remove exact duplicates, then rows with a value that is empty or
    #    couldn't be parsed, then any product (UPC) that appears twice
    df = drop_rows(df, issues, df.duplicated(), "Exact duplicate rows removed")
    missing = df.isna() | df.eq("")
    df = drop_rows(df, issues, missing.any(axis=1), "Rows with missing/unparseable values removed")
    df = drop_rows(df, issues, df.duplicated(subset="upc"), "Duplicate UPCs removed")

    # 5. Final types and a logical column order
    df = df.astype({"star_rating": "int64", "availability": "int64"})[COLUMN_ORDER]

    failed = [check for check, passed in validate(df) if not passed]
    if failed:
        raise ValueError(f"Clean data failed validation: {failed}")
    log.info("All validation checks passed")
    return df, issues


# --- Outputs ---

def write_excel(df: pd.DataFrame, path: Path) -> None:
    """Write a client-friendly Excel file: styled header, frozen top row,
    autofilter and column widths that fit the content."""
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Books")
        sheet = writer.sheets["Books"]

        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")

        sheet.freeze_panes = "A2"  # header stays visible while scrolling
        sheet.auto_filter.ref = sheet.dimensions  # filter dropdown on every column

        # Width = longest value in the column, capped so long titles/URLs don't take over
        # (headers get extra room for the filter arrow)
        for i, column in enumerate(df.columns, start=1):
            longest = max(df[column].astype(str).str.len().max(), len(column) + 3)
            sheet.column_dimensions[get_column_letter(i)].width = min(longest + 2, 60)

        price_letter = get_column_letter(df.columns.get_loc("price") + 1)
        for cell in sheet[price_letter][1:]:
            cell.number_format = "0.00"


def markdown_table(df: pd.DataFrame, raw: bool = False) -> list[str]:
    """Render a small DataFrame as Markdown table lines."""
    def cell(value) -> str:
        if raw:
            return as_code(value, max_length=100)
        if isinstance(value, float):
            return f"{value:.2f}"
        return str(value).replace("|", "\\|")

    lines = ["| " + " | ".join(df.columns) + " |", "|" + "---|" * len(df.columns)]
    lines += ["| " + " | ".join(cell(v) for v in row) + " |" for row in df.itertuples(index=False)]
    return lines


def pick_examples(raw: pd.DataFrame, clean: pd.DataFrame, n: int = 5) -> list:
    """Rows for the before/after table: the first few books plus books
    whose titles had broken characters (the most visible fix)."""
    broken_titles = raw.index[raw["Title"] != raw["Title"].map(fix_encoding)]
    candidates = list(clean.index[:3]) + list(broken_titles)
    return [i for i in dict.fromkeys(candidates) if i in clean.index][:n]


def write_report(raw: pd.DataFrame, clean: pd.DataFrame, issues: list[dict], path: Path) -> None:
    """Write a short Markdown report: row counts, issues fixed, checks and examples."""
    lines = [
        "# Cleaning report",
        "",
        f"Generated by `clean.py` on {date.today()} from `data/raw_books.csv`.",
        "",
        "## Row counts",
        "",
        "| | Rows |",
        "|---|---:|",
        f"| Before (`data/raw_books.csv`) | {len(raw)} |",
        f"| After (`data/clean_books.csv`) | {len(clean)} |",
        f"| Removed | {len(raw) - len(clean)} |",
        "",
        "## Issues found and fixed",
        "",
        "| Issue | Rows affected | Example (raw → clean) |",
        "|---|---:|---|",
    ]
    lines += [f"| {i['issue']} | {i['rows']} | {i['example']} |" for i in issues]

    lines += ["", "## Validation checks", "", "| Check | Result |", "|---|---|"]
    lines += [f"| {check} | {'✅ passed' if passed else '❌ failed'} |" for check, passed in validate(clean)]

    lines += ["", "## Column names", "", "| Raw | Clean |", "|---|---|"]
    lines += [f"| {name} | `{to_snake_case(name)}` |" for name in raw.columns]
    lines += ["| *(new)* | `currency` (split out of the price) |"]

    # Same title, different UPC: not a duplicate, but worth flagging to the client
    shared_titles = clean[clean["title"].duplicated(keep=False)]
    if not shared_titles.empty:
        lines += [
            "",
            "## Worth knowing",
            "",
            f"{shared_titles['title'].nunique()} title(s) appear more than once with different UPCs, prices "
            "and product pages. They are separate listings on the site, so all rows are kept:",
            "",
        ]
        lines += markdown_table(shared_titles[["title", "upc", "price", "availability"]])

    rows = pick_examples(raw, clean)
    lines += [
        "",
        f"## Before vs after: {len(rows)} example rows",
        "",
        "Raw values are shown in Python notation so hidden whitespace (`\\n`) and "
        "broken characters (`\\x80`) are visible. Product URLs are left out for width.",
        "",
        "**Before** (`data/raw_books.csv`)",
        "",
    ]
    lines += markdown_table(raw.loc[rows, ["UPC", "Title", "Price", "Star Rating", "Availability", "Category"]], raw=True)
    lines += ["", "**After** (`data/clean_books.csv`)", ""]
    lines += markdown_table(clean.loc[rows, ["upc", "title", "price", "currency", "star_rating", "availability", "category"]])

    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def run(raw_csv: Path = RAW_CSV) -> pd.DataFrame:
    """Clean the raw CSV and write the clean CSV, Excel file and report."""
    # Read every value as plain text and keep empty cells as "" (not NaN),
    # so the cleaning functions only ever see strings
    raw = pd.read_csv(raw_csv, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    log.info("Loaded %d raw rows from %s", len(raw), raw_csv.name)

    clean, issues = clean_books(raw)

    clean.to_csv(CLEAN_CSV, index=False, encoding="utf-8-sig")
    write_excel(clean, CLEAN_XLSX)
    write_report(raw, clean, issues, REPORT_MD)
    log.info("Saved %d clean rows to %s, %s and %s", len(clean), CLEAN_CSV.name, CLEAN_XLSX.name, REPORT_MD.name)
    return clean


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)-7s [%(name)s] %(message)s", datefmt="%H:%M:%S")
    run()
